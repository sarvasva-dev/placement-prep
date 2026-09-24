import os
import sys
import json
import docx
from docx.shared import Inches, Pt, RGBColor

HANDBOOK_SRC = r"D:\Projects\Placement_Master_Handbook\src"
if HANDBOOK_SRC not in sys.path:
    sys.path.insert(0, HANDBOOK_SRC)

from styles import (
    init_document, add_title, add_heading_1, add_heading_2, 
    add_heading_3, add_paragraph, add_bullet, add_callout, add_table, 
    add_code_block, add_checklist, add_tracker_table
)

WEB_ROOT = r"D:\Projects\Placement_Master_Handbook_Web"
CONTENT_DIR = os.path.join(WEB_ROOT, "content")
DAYS_DIR = os.path.join(CONTENT_DIR, "days")
EXPORTS_DIR = os.path.join(WEB_ROOT, "exports", "docx")
os.makedirs(EXPORTS_DIR, exist_ok=True)

OUTPUT_DOCX = os.path.join(EXPORTS_DIR, "Sarthak_30_Day_Placement_Master_Handbook.docx")

def build_master_docx():
    print("==========================================================")
    print("COMPILING 30-DAY MASTER HANDBOOK DOCX FROM CANONICAL STREAMS")
    print("==========================================================")
    
    doc = init_document()

    # Title Page
    add_title(
        doc,
        main_title="30-DAY PLACEMENT &\nSEMESTER MASTER HANDBOOK",
        subtitle="Complete Self-Contained Textbook, Solved PYQs, Aptitude Engine, Java 17+ DSA & Project Defense System",
        author="Sarthak Srivastava (3rd-Year BCA, CSJM University)",
        metadata={
            "Standard": "SGPA >= 9.0 Standard",
            "Edition": "September 2026 Canonical Master Edition",
            "Target": "Campus Placement Clearance & Production Readiness"
        }
    )

    doc.add_page_break()

    # Foreword & Architecture
    add_heading_1(doc, "Master Handbook Architecture & Ground Truth")
    add_paragraph(doc, "This master handbook is the offline compilation of the Sarthak 30-Day Placement & Semester Master Web System. It is generated directly from the identical canonical streams layer that powers the interactive study application.", bold_prefix="Single Source of Truth:")
    add_paragraph(doc, "Every day from Day 1 to Day 30 is a complete, self-contained study chapter containing: 1. Full textbook academic lectures with vector architecture models, 2. Verified university PYQ 15-mark model answers, 3. Authentic topic-specific worked aptitude problems with speed drills, 4. Java 17+ DSA solutions with execution traces, 5. Core CS lectures with interview Q&As, 6. Real project defenses grounded in production repository code, and 7. Daily 20-question mixed mastery tests.")

    # Loop through all 30 days
    for day_num in range(1, 31):
        day_path = os.path.join(DAYS_DIR, f"day{day_num:02d}.json")
        if not os.path.exists(day_path):
            print(f"[-] Missing day{day_num:02d}.json")
            continue
            
        with open(day_path, "r", encoding="utf-8") as f:
            d = json.load(f)

        doc.add_page_break()
        print(f"--> Rendering Day {day_num:02d}: {d.get('title', '')}")
        
        add_heading_1(doc, f"DAY {day_num:02d}: {d.get('title', '').upper()}")
        
        streams = d.get("streams", {})
        acad = streams.get("academic", {})
        pyqs = streams.get("academic_pyqs", [])
        apt_lesson = streams.get("aptitude_lesson", {})
        apt_solved = streams.get("aptitude_solved", [])
        dsa_pattern = streams.get("dsa_pattern", {})
        coding_probs = streams.get("coding_problems", [])
        cs = streams.get("core_cs", {})
        proj = streams.get("project_preparation", {})
        interviews = streams.get("placement_interview", [])
        mixed_test = streams.get("mixed_test", [])

        # Header Info
        subj_name = acad.get('subject_name') or acad.get('subject_code') or 'Academic'
        apt_topic = apt_lesson.get('topic') or 'Aptitude'
        add_paragraph(doc, f"Day {day_num} of 30  |  Subject Focus: {subj_name}  |  Aptitude: {apt_topic}  |  Target: SGPA >= 9.0 Standard", italic=True)

        # -------------------------------------------------------------
        # 1. SEMESTER ACADEMIC STUDY
        # -------------------------------------------------------------
        add_heading_2(doc, f"1. Semester 5 Deep Academic Lecture: {acad.get('topic', subj_name)}")
        
        if acad.get("objectives"):
            add_paragraph(doc, "Learning Objectives:", bold_prefix="Objectives:")
            for obj in acad["objectives"]:
                add_bullet(doc, obj)

        # Notes / Explanation
        explanation = acad.get("explanation")
        if isinstance(explanation, list):
            for p in explanation:
                add_paragraph(doc, p)
        elif isinstance(explanation, str) and explanation:
            for p in explanation.split("\n\n"):
                if p.strip(): add_paragraph(doc, p.strip())

        # Subtopics
        for sub in acad.get("subtopics", []):
            add_heading_3(doc, sub.get("title", ""))
            add_paragraph(doc, sub.get("content", ""))

        # System Architecture Model (Clean Callout, NO raw ASCII!)
        if acad.get("diagram"):
            add_callout(doc, f"ARCHITECTURAL FRAMEWORK & CONCEPTUAL FLOW: {acad.get('topic', '')}", [
                ("Interactive Vector Model:", f"Refer to the web visualizer for the full SVG rendering of {acad.get('topic')}.")
            ], box_type="understand")

        # Comparison Table
        if acad.get("comparison_table"):
            tbl = acad["comparison_table"]
            headers = tbl.get("headers", [])
            rows = tbl.get("rows", [])
            if headers and rows:
                widths = [Inches(6.5 / len(headers))] * len(headers)
                add_paragraph(doc, "Key Comparative Dimensions:", bold_prefix="Comparative Analysis:")
                add_table(doc, headers, rows, widths)

        # High-Yield Worked Numerical Problems (5004 Expansion)
        worked_nm = acad.get("worked_numerical_problems", [])
        if worked_nm:
            add_heading_3(doc, f"High-Yield Worked Numerical Problems ({len(worked_nm)} Fully Solved)")
            for widx, wp in enumerate(worked_nm, 1):
                add_paragraph(doc, f"Problem {widx}: {wp.get('problem', '')} [{wp.get('difficulty', 'Exam-Level')}]", bold_prefix=f"Problem {widx}:")
                add_paragraph(doc, wp.get("formula", ""), bold_prefix="Governing Formula:")
                for st in wp.get("step_by_step_calculation", []):
                    add_bullet(doc, st)
                
                # Iteration table if present
                iter_tbl = wp.get("iteration_table", [])
                if iter_tbl and isinstance(iter_tbl, list) and len(iter_tbl) > 0 and isinstance(iter_tbl[0], dict):
                    t_headers = [k.upper() for k in iter_tbl[0].keys()]
                    t_rows = [[str(v) for v in r.values()] for r in iter_tbl]
                    t_widths = [Inches(6.5 / len(t_headers))] * len(t_headers)
                    add_table(doc, t_headers, t_rows, t_widths)
                    
                add_paragraph(doc, f"Final Answer: {wp.get('final_answer', '')}  |  Verification: {wp.get('verification', '')}", bold_prefix="Result:")
                if wp.get("common_mistake"):
                    add_paragraph(doc, wp.get("common_mistake"), bold_prefix="Watch Out:")

        if acad.get("timed_numerical_drill"):
            td = acad["timed_numerical_drill"]
            add_callout(doc, f"TIMED NUMERICAL EXAM DRILL ({td.get('time_limit_minutes', 15)} Mins | {td.get('marks', 15)} Marks)", [
                ("Exam Problem:", td.get("problem", "")),
                ("Governing Formula:", td.get("formula", "")),
                ("Verified Answer:", td.get("final_answer", ""))
            ], box_type="exam")

        # Memorize / Understand / Traps
        if acad.get("memorize"):
            add_callout(doc, f"WHAT TO MEMORIZE (DAY {day_num})", [
                ("Core Definitions & Terminology:", acad['memorize'])
            ], box_type="memorize")

        if acad.get("understand"):
            add_callout(doc, f"WHAT TO UNDERSTAND DEEPLY (DAY {day_num})", [
                ("Conceptual Mechanics & Intuition:", acad['understand'])
            ], box_type="understand")

        if acad.get("common_mistakes"):
            add_callout(doc, f"EXAM MISTAKES TO AVOID (DAY {day_num})", [
                ("Common Mark-Losing Errors:", acad['common_mistakes'])
            ], box_type="trap")

        # -------------------------------------------------------------
        # 2. UNIVERSITY PYQ MODEL ANSWERS
        # -------------------------------------------------------------
        if pyqs:
            add_heading_2(doc, "2. CSJM University PYQs & Model Answers (15-Mark Standard)")
            for pidx, pyq in enumerate(pyqs, 1):
                add_callout(doc, f"UNIVERSITY PYQ #{pidx}: {pyq.get('year', 'CSJM University')} ({pyq.get('marks', 15)} Marks)", [
                    ("Exam Question:", pyq.get('question', '')),
                    ("Examiner Scoring Rubric:", pyq.get('rubric', 'Definition (3m) + Diagram (4m) + Technical Depth (5m) + Summary (3m) = 15 Marks'))
                ], box_type="pyq")

                if pyq.get("model_answer"):
                    add_heading_3(doc, f"Model Answer #{pidx} (15/15 Presentation)")
                    ma = pyq["model_answer"]
                    if isinstance(ma, list):
                        for item in ma:
                            if isinstance(item, list) and len(item) == 2:
                                add_paragraph(doc, item[1], bold_prefix=f"{item[0]}:")
                            elif isinstance(item, dict):
                                add_paragraph(doc, item.get('content', ''), bold_prefix=f"{item.get('heading', 'Key Point')}:")
                            else:
                                add_paragraph(doc, str(item))
                    elif isinstance(ma, str):
                        for p in ma.split("\n\n"):
                            if p.strip(): add_paragraph(doc, p.strip())

        # -------------------------------------------------------------
        # 3. APTITUDE (100% Topic-Specific Authentic Examples!)
        # -------------------------------------------------------------
        add_heading_2(doc, f"3. Quantitative & Placement Aptitude: {apt_topic}")
        if apt_lesson.get("concept"):
            add_paragraph(doc, apt_lesson.get("concept"))
            
        if apt_lesson.get("formulas"):
            add_callout(doc, "ESSENTIAL FORMULAS & SHORTCUTS", [
                ("Governing Formulas:", str(apt_lesson.get('formulas', ''))),
                ("Speed Shortcut:", str(apt_lesson.get('shortcuts', '')))
            ], box_type="shortcut")

        if apt_solved:
            add_heading_3(doc, f"Progressively Challenging Solved Problems ({len(apt_solved)} Solved)")
            for sidx, sp in enumerate(apt_solved, 1):
                tier_label = sp.get('tier') or f"Tier {sidx}"
                add_paragraph(doc, sp.get('question', sp.get('problem', '')), bold_prefix=f"Problem {sidx} ({tier_label}):")
                for st in sp.get('step_by_step_solution', sp.get('solution', [])):
                    add_bullet(doc, st)
                add_paragraph(doc, f"Answer: {sp.get('final_answer', sp.get('answer', ''))}  |  Shortcut: {sp.get('shortcut', '')}", bold_prefix="Result:")

        # -------------------------------------------------------------
        # 4. JAVA 17+ DSA PATTERN & CODING PROBLEMS
        # -------------------------------------------------------------
        if dsa_pattern and dsa_pattern.get("pattern_name"):
            add_heading_2(doc, f"4. Java 17+ DSA Pattern: {dsa_pattern.get('pattern_name', '')}")
            add_paragraph(doc, dsa_pattern.get('concept', ''), bold_prefix="Core Intuition:")
            java_code = dsa_pattern.get("java_code") or dsa_pattern.get("code", "")
            if java_code:
                add_code_block(doc, java_code, language="java")
            add_paragraph(doc, f"Time: {dsa_pattern.get('time_complexity', 'O(N)')}  |  Space: {dsa_pattern.get('space_complexity', 'O(1)')}", bold_prefix="Complexity:")

        if coding_probs:
            add_heading_3(doc, "High-Frequency Java 17+ Placement Coding Problems")
            for cp in coding_probs:
                add_paragraph(doc, cp.get('problem_statement', cp.get('problem', '')), bold_prefix=f"Problem: {cp.get('title', '')}")
                add_paragraph(doc, cp.get('solution_approach', cp.get('approach', '')), bold_prefix="Algorithmic Strategy:")
                c_code = cp.get("solution_java") or cp.get("code", "")
                if c_code:
                    add_code_block(doc, c_code, language="java")
                add_paragraph(doc, f"Time: {cp.get('time_complexity', 'O(N)')}  |  Space: {cp.get('space_complexity', 'O(1)')}", bold_prefix="Complexities:")

        # -------------------------------------------------------------
        # 5. CORE COMPUTER SCIENCE
        # -------------------------------------------------------------
        if cs and cs.get("topic"):
            add_heading_2(doc, f"5. Core Computer Science: {cs.get('subject', '')} — {cs.get('topic', '')}")
            for note in cs.get("detailed_notes", []):
                add_paragraph(doc, note)
            for qa in cs.get("interview_qa", []):
                add_callout(doc, f"CORE CS INTERVIEW Q&A: {qa.get('q', '')}", [
                    ("Expected Technical Answer:", qa.get('a', ''))
                ], box_type="understand")

        # -------------------------------------------------------------
        # 6. PROJECT DEFENSE & PLACEMENT INTERVIEW
        # -------------------------------------------------------------
        if proj and proj.get("project_name"):
            add_heading_2(doc, f"6. Project Architecture Defense: {proj.get('project_name', '')}")
            add_paragraph(doc, proj.get('architecture_deep_dive', ''))
            for qa in proj.get("interview_qa", []):
                add_callout(doc, f"PROJECT INTERVIEW DEFENSE: {qa.get('q', '')}", [
                    ("Architectural Defense:", qa.get('a', ''))
                ], box_type="memorize")

        if interviews:
            add_heading_3(doc, "Placement HR & Technical STAR Behavioral Answers")
            for iv in interviews:
                add_callout(doc, f"INTERVIEW QUESTION: {iv.get('question', '')}", [
                    ("Strong Candidate Response (STAR Framework):", iv.get('answer', ''))
                ], box_type="interview")

        # -------------------------------------------------------------
        # 7. DAILY 20-QUESTION MIXED MASTERY TEST
        # -------------------------------------------------------------
        if mixed_test:
            add_heading_2(doc, f"7. Day {day_num:02d} Comprehensive 20-Question Daily Mastery Test")
            add_paragraph(doc, f"Diagnostic evaluation covering Academic ({subj_name}), Aptitude ({apt_topic}), Java DSA, and Core CS. Total: {len(mixed_test)} Questions.", italic=True)
            
            test_rows = []
            for idx, q in enumerate(mixed_test, 1):
                opts = q.get("options", {})
                if isinstance(opts, dict):
                    opt_str = f"A: {opts.get('A','')} | B: {opts.get('B','')} | C: {opts.get('C','')} | D: {opts.get('D','')}"
                elif isinstance(opts, list):
                    opt_str = " | ".join(opts)
                else:
                    opt_str = ""
                    
                q_text = f"{q.get('question', '')}\nOptions: {opt_str}"
                ans_text = f"Ans: {q.get('correct_answer', '')}\n{q.get('explanation', '')}"
                test_rows.append([f"Q{idx}", q_text, ans_text])
                
            add_table(doc, ["No.", "Question & Multiple Choice Options", "Correct Answer & Technical Rationale"], test_rows, [Inches(0.5), Inches(3.4), Inches(2.6)])

    # Save final DOCX
    print(f"\n--> Saving final DOCX to: {OUTPUT_DOCX}")
    doc.save(OUTPUT_DOCX)
    print(f"--> SUCCESS! Master Handbook DOCX compiled ({os.path.getsize(OUTPUT_DOCX) / 1024 / 1024:.2f} MB)")

if __name__ == "__main__":
    build_master_docx()
